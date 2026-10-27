#!/usr/bin/env python3
"""Genera un consultor HTML standalone sobre la hoja Subcontratos del master MAST.

Responde, para cualquier proveedor: cuántos pagos, cuánto, en qué OTs, y qué
estados tienen. Incluye el detalle fila por fila y los hallazgos del validador
que tocan a ese proveedor.

Los datos van embebidos en el HTML (sin servidor, se abre con doble clic).
NO se publica: contiene montos y RUT de proveedores.

Uso:
    python 06_consultor_subcontratos.py
    open -a "Google Chrome" consultor-subcontratos.html
"""
import json
import os
from collections import Counter, defaultdict

import openpyxl

XLSX = "/Users/victorvilche/Documents/EjercicioJose/MASTER - Venta-Sub-GM_Flujos.xlsx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "consultor-subcontratos.html")
HALLAZGOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "hallazgos-validador.json")


def extraer():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    wbv = openpyxl.load_workbook(XLSX)
    ws, wsv = wb["Subcontratos"], wbv["Subcontratos"]
    hdr = {}
    for c in wsv[1]:
        if c.value is not None:
            hdr[str(c.value).strip()] = c.column

    def g(r, name):
        c = hdr.get(name)
        return ws.cell(r, c).value if c else None

    filas = []
    for r in range(2, ws.max_row + 1):
        prov = g(r, "Proveedor")
        if not prov or str(prov).strip() == "":
            continue
        if str(prov).strip().upper().startswith(("TOTAL", "SUBTOTAL")):
            continue
        bruto = g(r, "Pago Bruto")
        neto = g(r, "Pago Neto")
        filas.append({
            "fila": r,
            "centro": str(g(r, "Centro de Costos") or "").strip(),
            "proyecto": str(g(r, "Proyecto") or "").strip(),
            "tipo_orden": str(g(r, "Tipo Orden") or "").strip(),
            "ot": str(g(r, "N° OT") or "").strip(),
            "proveedor": str(prov).strip(),
            "rut": str(g(r, "Rut") or "").strip(),
            "concepto": str(g(r, "Concepto") or "").strip(),
            "tipo_doc": str(g(r, "Tipo Documento") or "").strip(),
            "factura": str(g(r, "N° de factura") or "").strip(),
            "estado_fact": str(g(r, "Estado facturación") or "").strip(),
            "estado_pago": str(g(r, "Estado del pago") or "").strip(),
            "hito": str(g(r, "Hito de pago") or "").strip(),
            "neto": round(float(neto), 2) if isinstance(neto, (int, float)) else None,
            "bruto": round(float(bruto), 2) if isinstance(bruto, (int, float)) else None,
            "pct": float(g(r, "% Pago")) if isinstance(g(r, "% Pago"), (int, float)) else None,
            "total_neto": round(float(g(r, "TOTAL NETO")), 2) if isinstance(g(r, "TOTAL NETO"), (int, float)) else None,
        })

    # hallazgos por fila, si el validador ya corrió
    por_fila = defaultdict(list)
    if os.path.exists(HALLAZGOS):
        try:
            for h in json.load(open(HALLAZGOS, encoding="utf-8"))["hallazgos"]:
                if h["hoja"] == "Subcontratos" and h["fila"]:
                    por_fila[h["fila"]].append(f"{h['codigo']} ({h['severidad']})")
        except (json.JSONDecodeError, KeyError):
            pass

    for f in filas:
        f["alertas"] = por_fila.get(f["fila"], [])

    return filas


def agrupar(filas):
    """Agrupa por RUT cuando existe: el mismo proveedor aparece escrito de 2-3 formas.
    Sin RUT, cae al nombre normalizado (mayúsculas, espacios colapsados, sin acentos)."""
    import unicodedata

    def norm(s):
        s = unicodedata.normalize("NFKD", s)
        s = "".join(c for c in s if not unicodedata.combining(c))
        return " ".join(s.split()).upper()

    provs = defaultdict(lambda: {"filas": [], "pagos": 0, "bruto": 0.0,
                                 "pagado": 0.0, "sin_pago": 0.0, "eliminada": 0.0,
                                 "ots": set(), "proyectos": set(), "alertas": 0,
                                 "nombres": set(), "ruts": set()})
    for f in filas:
        clave = ("RUT", f["rut"]) if f["rut"] else ("NOM", norm(f["proveedor"]))
        p = provs[clave]
        p["filas"].append(f)
        p["pagos"] += 1
        b = f["bruto"] or 0.0
        p["bruto"] += b
        e = f["estado_pago"].lower()
        if e == "pagado":
            p["pagado"] += b
        elif e == "sin pago":
            p["sin_pago"] += b
        elif e in ("eliminada", "eliminado"):
            p["eliminada"] += b
        if f["ot"]:
            p["ots"].add(f["ot"])
        if f["proyecto"]:
            p["proyectos"].add(f["proyecto"])
        p["alertas"] += len(f["alertas"])
        p["nombres"].add(" ".join(f["proveedor"].split()))
        if f["rut"]:
            p["ruts"].add(f["rut"])

    salida = []
    for clave, p in provs.items():
        nombres = sorted(p["nombres"], key=len, reverse=True)
        salida.append({
            "proveedor": nombres[0],
            "variantes": nombres[1:],
            "rut": next(iter(p["ruts"]), ""),
            "pagos": p["pagos"],
            "bruto": round(p["bruto"], 2),
            "pagado": round(p["pagado"], 2),
            "sinPago": round(p["sin_pago"], 2),
            "eliminada": round(p["eliminada"], 2),
            "n_ots": len(p["ots"]),
            "n_proyectos": len(p["proyectos"]),
            "alertas": p["alertas"],
        })
    salida.sort(key=lambda x: -x["bruto"])
    return salida


HTML = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Consultor de Subcontratos — MAST</title>
<style>
  :root{--bg:#0e1116;--panel:#161b22;--panel2:#1c232c;--line:#2d3742;--txt:#e6edf3;
        --dim:#8b949e;--acc:#4da3ff;--ok:#3fb950;--warn:#d29922;--bad:#f85149;
        --mono:ui-monospace,SFMono-Regular,Menlo,monospace}
  *{box-sizing:border-box;margin:0;padding:0}
  body{background:var(--bg);color:var(--txt);font:14.5px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;padding:26px 22px 60px;max-width:1320px;margin:0 auto}
  h1{font-size:23px;font-weight:650;letter-spacing:-.02em}
  .sub{color:var(--dim);font-size:13px;margin-top:5px}
  .eyebrow{font:600 10.5px/1 var(--mono);letter-spacing:.12em;text-transform:uppercase;color:var(--dim);margin-bottom:7px}
  .grid{display:grid;gap:10px}
  .kpis{grid-template-columns:repeat(5,1fr);margin:20px 0}
  @media(max-width:900px){.kpis{grid-template-columns:repeat(2,1fr)}}
  .kpi{background:var(--panel);border:1px solid var(--line);border-radius:9px;padding:11px 13px}
  .kpi b{display:block;font:650 20px/1.15 var(--mono);color:var(--acc)}
  .kpi span{font-size:11.5px;color:var(--dim);display:block;margin-top:2px}
  input[type=search]{width:100%;background:var(--panel);border:1px solid var(--line);color:var(--txt);
    padding:11px 14px;border-radius:8px;font-size:14.5px;outline:none}
  input[type=search]:focus{border-color:var(--acc)}
  .layout{display:grid;grid-template-columns:330px 1fr;gap:14px;margin-top:14px;align-items:start}
  @media(max-width:900px){.layout{grid-template-columns:1fr}}
  .lista{background:var(--panel);border:1px solid var(--line);border-radius:9px;max-height:620px;overflow-y:auto}
  .item{padding:9px 12px;border-bottom:1px solid #21262d;cursor:pointer;font-size:13px}
  .item:hover{background:var(--panel2)}
  .item.sel{background:#1f3a5f;border-left:3px solid var(--acc)}
  .item b{display:block;font-size:13px;font-weight:600}
  .item .m{font:11.5px var(--mono);color:var(--dim)}
  .panel{background:var(--panel);border:1px solid var(--line);border-radius:9px;padding:16px 18px}
  table{width:100%;border-collapse:collapse;font-size:12.5px;margin-top:10px}
  th{text-align:left;padding:7px 8px;border-bottom:1px solid var(--line);color:var(--dim);
     font:600 10px/1 var(--mono);letter-spacing:.07em;text-transform:uppercase;position:sticky;top:0;background:var(--panel)}
  td{padding:6px 8px;border-bottom:1px solid #21262d;vertical-align:top}
  tr:hover td{background:var(--panel2)}
  .num{font-family:var(--mono);text-align:right;white-space:nowrap}
  .tag{font:600 9.5px/1 var(--mono);padding:3px 6px;border-radius:4px;text-transform:uppercase;white-space:nowrap}
  .t-pag{background:#1a3d24;color:#56d364}
  .t-sin{background:#3d2a1a;color:#d29922}
  .t-eli{background:#3d1f24;color:#ff7b72}
  .t-x{background:#21262d;color:#8b949e}
  .alert{background:#3d1f24;color:#ff9b95;font:600 9.5px var(--mono);padding:3px 6px;border-radius:4px}
  .empty{color:var(--dim);padding:26px;text-align:center;font-size:13px}
  .bar{height:5px;background:#21262d;border-radius:3px;overflow:hidden;margin-top:7px}
  .bar i{display:block;height:100%;background:var(--acc)}
  .big{font:650 26px/1.1 var(--mono);color:var(--txt)}
  .row{display:flex;gap:22px;flex-wrap:wrap;margin:12px 0 4px}
  .row div{min-width:110px}
  .row small{display:block;font:600 9.5px var(--mono);color:var(--dim);text-transform:uppercase;letter-spacing:.07em}
  .scroll{max-height:430px;overflow-y:auto;margin-top:10px}
</style>
</head>
<body>
<div class="eyebrow">MAST Arquitectura y Construcción · hoja Subcontratos del master</div>
<h1>Consultor de Subcontratos</h1>
<p class="sub" id="fuente"></p>

<div class="grid kpis" id="kpis"></div>

<input type="search" id="buscar" placeholder="Buscar proveedor, RUT, proyecto, OT o concepto…" autocomplete="off">

<div class="layout">
  <div class="lista" id="lista"></div>
  <div class="panel" id="detalle"><div class="empty">Selecciona un proveedor de la izquierda</div></div>
</div>

<script>
const DATA = __DATA__;
const PROV = __PROV__;
const MON = n => (n==null? "—" : "$" + Math.round(n).toLocaleString("es-CL"));
const P = n => (n==null? "—" : (n*100).toFixed(1) + "%");
const esc = s => String(s==null?"":s).replace(/[&<>"]/g, c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));

document.getElementById("fuente").textContent =
  DATA.length + " pagos en la hoja · " + PROV.length + " proveedores · fuente: MASTER - Venta-Sub-GM_Flujos.xlsx";

const totPagos = DATA.length;
const totBruto = DATA.reduce((a,f)=>a+(f.bruto||0),0);
const conAlerta = DATA.filter(f=>f.alertas.length).length;
document.getElementById("kpis").innerHTML = [
  ["Pagos registrados", totPagos.toLocaleString("es-CL")],
  ["Proveedores", PROV.length],
  ["Monto total bruto", MON(totBruto)],
  ["OTs distintas", new Set(DATA.map(f=>f.ot).filter(Boolean)).size],
  ["Pagos con alerta", conAlerta],
].map(([k,v])=>`<div class="kpi"><b>${v}</b><span>${k}</span></div>`).join("");

let sel = null, filtro = "";

function filtrados(){
  const q = filtro.trim().toLowerCase();
  if(!q) return PROV;
  return PROV.filter(p => {
    if (p.proveedor.toLowerCase().includes(q)) return true;
    if ((p.rut||"").toLowerCase().includes(q)) return true;
    if ((p.variantes||[]).some(v => v.toLowerCase().includes(q))) return true;
    // buscar también dentro de las filas del proveedor (proyecto, OT, concepto, hito)
    const nombres = new Set([p.proveedor, ...(p.variantes||[])]);
    return DATA.some(f => (p.rut ? f.rut === p.rut : nombres.has(f.proveedor))
      && (f.proyecto+" "+f.ot+" "+f.concepto+" "+f.hito).toLowerCase().includes(q));
  });
}

function pintarLista(){
  const xs = filtrados();
  const max = Math.max(...xs.map(p=>p.bruto), 1);
  document.getElementById("lista").innerHTML = xs.length ? xs.map(p=>`
    <div class="item ${sel===p.proveedor?'sel':''}" onclick="elegir('${esc(p.proveedor).replace(/'/g,"\\\\'")}')">
      <b>${esc(p.proveedor)}</b>
      <div class="m">${p.pagos} pago${p.pagos!==1?'s':''} · ${p.n_ots} OT · ${MON(p.bruto)}</div>
      ${p.alertas?`<span class="alert">${p.alertas} alerta${p.alertas!==1?'s':''}</span>`:""}
      <div class="bar"><i style="width:${(p.bruto/max*100).toFixed(1)}%"></i></div>
    </div>`).join("") : '<div class="empty">Sin resultados</div>';
}

function elegir(nombre){
  sel = nombre;
  pintarLista();
  const p = PROV.find(x=>x.proveedor===nombre);
  // las filas NO vienen en PROV (se quitaron para no duplicar 663 registros);
  // se reconstruyen desde DATA filtrando por RUT o por nombre/variantes.
  const nombres = new Set([p.proveedor, ...(p.variantes||[])]);
  const filas = DATA.filter(f => p.rut ? f.rut === p.rut : nombres.has(f.proveedor))
                    .slice().sort((a,b)=>(a.ot||"").localeCompare(b.ot||"") || a.fila-b.fila);
  // los totales se recalculan desde las MISMAS filas que se muestran, redondeando
  // cada fila antes de sumar, para que el panel cuadre con la columna Bruto de la tabla.
  const r = x => Math.round(x || 0);
  const t = {bruto:0, pagado:0, sinPago:0, eliminada:0};
  filas.forEach(f => {
    const b = r(f.bruto);
    t.bruto += b;
    const e = (f.estado_pago||"").toLowerCase();
    if (e === "pagado") t.pagado += b;
    else if (e === "sin pago") t.sinPago += b;
    else if (e.startsWith("elimin")) t.eliminada += b;
  });
  const porEstado = {};
  filas.forEach(f=>{ porEstado[f.estado_pago] = (porEstado[f.estado_pago]||0) + 1; });
  const porOt = {};
  filas.forEach(f=>{ const k=f.ot||"(sin OT)"; porOt[k] = (porOt[k]||0) + 1; });

  document.getElementById("detalle").innerHTML = `
    <div class="eyebrow">${p.rut ? "RUT " + esc(p.rut) : "sin RUT registrado"}</div>
    <div class="big">${esc(p.proveedor)}</div>
    ${p.variantes && p.variantes.length ? `<p class="sub" style="color:var(--warn)">El mismo proveedor aparece escrito también como: ${p.variantes.map(esc).join(" · ")}</p>` : ""}

    <div class="row">
      <div><small>Pagos</small><span class="big">${filas.length}</span></div>
      <div><small>Total bruto</small><span style="font:600 17px var(--mono)">${MON(t.bruto)}</span></div>
      <div><small>Pagado</small><span style="font:600 17px var(--mono);color:var(--ok)">${MON(t.pagado)}</span></div>
      <div><small>Sin pago</small><span style="font:600 17px var(--mono);color:var(--warn)">${MON(t.sinPago)}</span></div>
      <div><small>Eliminado</small><span style="font:600 17px var(--mono);color:var(--bad)">${MON(t.eliminada)}</span></div>
      <div><small>OTs</small><span style="font:600 17px var(--mono)">${Object.keys(porOt).filter(k=>k!=="(sin OT)").length}</span></div>
      <div><small>Proyectos</small><span style="font:600 17px var(--mono)">${new Set(filas.map(f=>f.proyecto).filter(Boolean)).size}</span></div>
    </div>

    <p class="sub" style="margin-top:14px">Estados: ${Object.entries(porEstado).map(([k,v])=>`${esc(k)} (${v})`).join(" · ")}</p>

    <div class="scroll">
    <table>
      <thead><tr>
        <th>Fila</th><th>Tipo</th><th>OT</th><th>Hito</th><th>Doc</th>
        <th class="num">Total</th><th class="num">%</th><th class="num">Neto</th><th class="num">Bruto</th>
        <th>Estado</th><th>Proyecto</th>
      </tr></thead>
      <tbody>
      ${filas.map(f=>`
        <tr>
          <td><span class="m">${f.fila}</span></td>
          <td>${esc(f.tipo_orden)}</td>
          <td title="${esc(f.ot)}">${esc(f.ot.length>22?f.ot.slice(0,22)+"…":f.ot)}</td>
          <td>${esc(f.hito)}</td>
          <td>${esc(f.tipo_doc)}</td>
          <td class="num">${MON(f.total_neto)}</td>
          <td class="num">${P(f.pct)}</td>
          <td class="num">${MON(f.neto)}</td>
          <td class="num">${MON(f.bruto)}</td>
          <td><span class="tag ${f.estado_pago.toLowerCase()==='pagado'?'t-pag':f.estado_pago.toLowerCase()==='sin pago'?'t-sin':f.estado_pago.toLowerCase().startsWith('elimin')?'t-eli':'t-x'}">${esc(f.estado_pago||"—")}</span>
              ${f.alertas.length?`<br><span class="alert">${esc(f.alertas.join(", "))}</span>`:""}</td>
          <td>${esc(f.proyecto)}</td>
        </tr>`).join("")}
      </tbody>
    </table>
    </div>

    <p class="sub" style="margin-top:12px">Por OT: ${Object.entries(porOt).map(([k,v])=>`${esc(k)} (${v})`).join(" · ")}</p>
  `;
}

document.getElementById("buscar").addEventListener("input", e=>{ filtro=e.target.value; pintarLista(); });
pintarLista();
if(PROV.length) elegir(PROV[0].proveedor);
</script>
</body>
</html>
"""


def main():
    filas = extraer()
    provs = agrupar(filas)

    html = HTML.replace("__DATA__", json.dumps(filas, ensure_ascii=False)) \
               .replace("__PROV__", json.dumps(provs, ensure_ascii=False))
    open(OUT, "w", encoding="utf-8").write(html)

    con_variantes = [p for p in provs if p["variantes"]]
    print(f"pagos: {len(filas)}  proveedores (agrupados por RUT): {len(provs)}")
    print(f"bruto total: ${sum(f['bruto'] or 0 for f in filas):,.0f}")
    print(f"proveedores con nombre escrito de varias formas: {len(con_variantes)}")
    for p in con_variantes:
        print(f"  {p['proveedor'][:44]:44s} + {len(p['variantes'])} variante(s)")
    print(f"\nescrito: {OUT}")
    print("\nTop 5 por monto:")
    for p in provs[:5]:
        print(f"  {p['proveedor'][:38]:38s} {p['pagos']:3d} pagos  ${p['bruto']:>14,.0f}  {p['n_ots']:3d} OTs")


if __name__ == "__main__":
    main()
