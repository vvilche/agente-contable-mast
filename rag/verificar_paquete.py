#!/usr/bin/env python3
"""Verificación ad-hoc del paquete de plan + RAG (no es una suite del proyecto).

Comprueba que los documentos entregados sean coherentes CON LOS ARCHIVOS REALES
de José, no solo entre sí. Un JSON que se contradice con el Excel es peor que no
tener JSON: el agente heredaría el error.

Uso: ~/.hermes/hermes-agent/venv/bin/python hermes-verify-mast.py
     exit 0 = todo cuadra | exit 1 = hay desajustes
"""
import json
import os
import re
import sys
from collections import Counter

import openpyxl

BASE = "/Users/victorvilche/orca/workspaces/Clases Usach/EjercicioClases"
XLSX = "/Users/victorvilche/Documents/EjercicioJose/MASTER - Venta-Sub-GM_Flujos.xlsx"
PDF = "/Users/victorvilche/Documents/EjercicioJose/PROTOCOLO DE GESTIÓN INTERNA.pdf"

fails = []


def check(cond, msg):
    print(("  OK    " if cond else "  FALLA ") + msg)
    if not cond:
        fails.append(msg)


# ── 1. JSON válidos y forma de los 4 dominios ────────────────────────────
print("\n[1] Manifiesto: 4 dominios")
man = json.load(open(f"{BASE}/rag/01-manifiesto-corpus.json", encoding="utf-8"))
check([d["id"] for d in man["dominios"]] == ["A", "B", "C", "D"], "dominios A,B,C,D en orden")
for d in man["dominios"]:
    check(all(k in d for k in ("nombre", "tipo_indice")), f"dominio {d['id']} con nombre y tipo_indice")

d_dom = next(d for d in man["dominios"] if d["id"] == "D")
check("NONE" in d_dom["tipo_indice"], "dominio D es lookup determinista, no vectorial")
check("fuente_existente" in d_dom, "dominio D declara la fuente existente reutilizada")

# ── 2. Conteos del JSON vs Excel real ────────────────────────────────────
print("\n[2] Conteos declarados vs Excel medido")
wsv = openpyxl.load_workbook(XLSX, data_only=True)["Ventas global"]
wsf = openpyxl.load_workbook(XLSX)["Ventas global"]
hdr = [c.value for c in wsf[1]]
col = lambda name: [wsv.cell(r, hdr.index(name) + 1).value for r in range(2, wsv.max_row + 1)]

filas = sum(1 for r in range(2, wsv.max_row + 1) if any(wsv.cell(r, c).value is not None for c in range(1, 21)))
estado = Counter(x for x in col("Estado del pago") if x)
tipodoc = Counter(x for x in col("Tipo de documento") if x)
m = json.load(open(f"{BASE}/rag/02-reglas-mast.json", encoding="utf-8"))
med = m["valores_observados_que_el_agente_debe_validar"]["conteo_medido_ventas_global"]

check(med["filas_totales"] == filas, f"filas_totales={med['filas_totales']} vs medido {filas}")
for k, v in (("pagado", "Pagado"), ("factorizada", "Factorizada"), ("sin_pago", "Sin Pago"), ("eliminado", "Eliminado")):
    check(med[k] == estado[v], f"{k}={med[k]} vs medido {estado[v]}")
check(sum(estado.values()) == med["suma_por_estado_pago"], f"suma por estado {sum(estado.values())} = declarada {med['suma_por_estado_pago']}")
for k, v in med["por_tipo"].items():
    check(tipodoc[k] == v, f"tipo {k}={v} vs medido {tipodoc[k]}")

# ── 3. Las anomalías declaradas deben existir de verdad ─────────────────
print("\n[3] Anomalías: verificadas contra el Excel")
iN, iP, iQ, iT = (hdr.index(c) + 1 for c in ("TOTAL NETO", "Pago Neto", "Pago Bruto", "Tipo de documento"))
exentas = con_iva = 0
for r in range(2, wsv.max_row + 1):
    if wsv.cell(r, iT).value == "Factura Exenta":
        exentas += 1
        p, q = wsv.cell(r, iP).value, wsv.cell(r, iQ).value
        try:
            if p and q and round(float(q) / float(p), 4) == 1.19:
                con_iva += 1
        except (TypeError, ValueError):
            pass
check((exentas, con_iva) == (6, 3), f"ANOM-02: {exentas} exentas, {con_iva} con IVA (JSON: 6 y 3)")

rotas = [r for r in range(2, wsv.max_row + 1)
         if isinstance(wsv.cell(r, iN).value, (int, float)) and wsv.cell(r, iN).value < 1000]
check(len(rotas) == 14, f"ANOM-05: {len(rotas)} filas con TOTAL NETO < 1000 (JSON: 14)")

# ── 4. Reglas: forma y citas contra el PDF ──────────────────────────────
print("\n[4] Reglas y citas")
reglas = m["reglas"]
check(len({r["id"] for r in reglas}) == len(reglas), f"{len(reglas)} ids únicos")
check(all(all(k in r for k in ("id", "regla", "cita_literal", "pagina", "seccion", "verificado")) for r in reglas),
      "todas las reglas con id/regla/cita_literal/pagina/seccion/verificado")

try:
    import pymupdf
    doc = pymupdf.open(PDF)

    def norm(s):
        s = re.sub(r"[^a-z0-9%áéíóúñü ]", " ", s.lower())
        return re.sub(r"\s+", " ", s).strip()

    malas = []
    for r in reglas:
        pg = norm(doc[r["pagina"] - 1].get_text())
        for trozo in (t for t in r["cita_literal"].split(" / ") if len(t.strip()) > 10):
            if " ".join(norm(trozo).split()[:6]) not in pg:
                malas.append(r["id"])
                break
    check(not malas, f"{len(reglas) - len(malas)}/{len(reglas)} citas halladas en su página" + (f" — fallan {malas}" if malas else ""))
except ImportError:
    print("  SKIP  pymupdf no disponible: verificación de citas omitida")

# ── 5. HTML: estructura, sin placeholders reales, secciones numeradas ───
print("\n[5] HTML de arquitectura")
h = open(f"{BASE}/arquitectura-agente-mast.html", encoding="utf-8").read()
# `}}` es sintaxis CSS válida (@media anidado): solo los placeholders de plantilla son sospechosos
check(not re.search(r"\{\{[^}]*\}\}|\bTODO\b|\bFIXME\b", h), "sin placeholders de plantilla (se ignora `}}` de CSS)")
for tag in ("details", "summary", "table", "div"):
    a, c = len(re.findall(rf"<{tag}[ >]", h)), h.count(f"</{tag}>")
    check(a == c, f"<{tag}> balanceado ({a}/{c})")
check(all(f"D{n}</span>" in h for n in range(1, 8)), "subsistemas D1..D7 presentes")
nums = [int(x) for x in re.findall(r"<h2>(\d+) ·", h)]
check(nums == sorted(set(nums)) and len(nums) == len(set(nums)), f"secciones numeradas sin duplicados: {nums}")
check("ANOM" not in h and "4.778.748" in h, "el hallazgo de las 14 filas está en el HTML")

# ── Resumen ─────────────────────────────────────────────────────────────
print("\n" + "=" * 54)
if fails:
    print(f"{len(fails)} FALLA(S):")
    for f in fails:
        print("  -", f)
    sys.exit(1)
print("TODO OK")
