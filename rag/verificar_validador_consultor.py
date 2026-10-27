#!/usr/bin/env python3
"""Verificación ad-hoc del validador y del consultor de subcontratos.

No es una suite del proyecto (el repo son documentos y scripts de análisis).
Comprueba que las CIFRAS que muestran las herramientas cuadren contra el Excel real,
porque un consultor que muestra un número equivocado es peor que no tener consultor.

Uso: ~/.hermes/hermes-agent/venv/bin/python hermes-verify-validador-consultor.py
"""
import json
import os
import re
import subprocess
import sys
from collections import Counter, defaultdict

import openpyxl

RAG = "/Users/victorvilche/orca/workspaces/Clases Usach/EjercicioClases/rag"
XLSX = "/Users/victorvilche/Documents/EjercicioJose/MASTER - Venta-Sub-GM_Flujos.xlsx"
HTML = f"{RAG}/consultor-subcontratos.html"
OUTV = f"{RAG}/hallazgos-validador.json"
PY = sys.executable
fails = []


def check(c, m):
    print(("  OK    " if c else "  FALLA ") + m)
    if not c:
        fails.append(m)


# ── 1. Validador: corre y produce hallazgos coherentes ───────────────────
print("\n[1] Validador")
r = subprocess.run([PY, "05_validador_master.py", "--json", OUTV], cwd=RAG,
                   capture_output=True, text=True)
check(r.returncode in (0, 2), f"corre (exit {r.returncode})")
val = json.load(open(OUTV, encoding="utf-8"))
h = val["hallazgos"]
check(len(h) > 0, f"{len(h)} hallazgos emitidos")
codigos = Counter(x["codigo"] for x in h)
print("       " + ", ".join(f"{k}={v}" for k, v in codigos.most_common()))
check(codigos.get("PAGO_DUPLICADO", 0) == 4, "4 pagos duplicados detectados")
check(codigos.get("FILA_TOTAL_COMO_DATO", 0) == 1, "1 fila de total metida como dato")
check(all(x["regla"] for x in h), "cada hallazgo cita la regla que lo sostiene")

# ── 2. Los duplicados existen de verdad en el Excel ──────────────────────
print("\n[2] Duplicados verificados contra el Excel")
wb = openpyxl.load_workbook(XLSX, data_only=True)
wbv = openpyxl.load_workbook(XLSX)
ws, wsv = wb["Subcontratos"], wbv["Subcontratos"]
hdr = {str(c.value).strip(): c.column for c in wsv[1] if c.value is not None}
dup_val = [x for x in h if x["codigo"] == "PAGO_DUPLICADO"]
for x in dup_val[:2]:
    fila, prev = x["fila"], int(re.search(r"fila (\d+)", x["mensaje"]).group(1))
    iguales = all(ws.cell(fila, hdr[k]).value == ws.cell(prev, hdr[k]).value
                  for k in ("Proveedor", "N° OT", "Hito de pago", "Pago Bruto"))
    check(iguales, f"fila {fila} repite realmente a la fila {prev}")

# ── 3. Consultor: JSON embebido válido y sumas cuadran ───────────────────
print("\n[3] Consultor: datos embebidos")
html = open(HTML, encoding="utf-8").read()
m = re.search(r"const DATA = (\[.*?\]);\nconst PROV = (\[.*?\]);", html, re.S)
check(bool(m), "DATA y PROV presentes en el HTML")
DATA = json.loads(m.group(1))
PROV = json.loads(m.group(2))
check(len(DATA) == 663, f"663 pagos ({len(DATA)})")
check(len(PROV) == 89, f"89 proveedores agrupados por RUT ({len(PROV)})")

suma_data = sum(f["bruto"] or 0 for f in DATA)
suma_prov = sum(p["bruto"] for p in PROV)
check(abs(suma_data - suma_prov) < 1, f"suma por pago = suma por proveedor (${suma_data:,.0f})")

# contra el Excel: el bruto total debe coincidir
cQ = hdr["Pago Bruto"]
prov_col = hdr["Proveedor"]
excel_bruto = sum(ws.cell(r, cQ).value for r in range(2, ws.max_row + 1)
                  if isinstance(ws.cell(r, cQ).value, (int, float))
                  and ws.cell(r, prov_col).value
                  and not str(ws.cell(r, prov_col).value).upper().startswith("TOTAL"))
check(abs(excel_bruto - suma_data) < 1, f"bruto total coincide con el Excel (${excel_bruto:,.0f})")

# ── 4. Agrupación por RUT: un RUT no aparece en dos proveedores ──────────
print("\n[4] Un proveedor = un RUT")
ruts = [p["rut"] for p in PROV if p["rut"]]
check(len(ruts) == len(set(ruts)), f"sin RUT repetidos entre proveedores ({len(ruts)} con RUT)")
con_var = [p for p in PROV if p.get("variantes")]
check(len(con_var) == 7, f"7 proveedores con nombre escrito de varias formas ({len(con_var)})")
for p in con_var[:3]:
    print(f"       {p['proveedor'][:40]} -> {p['variantes']}")

# ── 5. El HTML no tiene los bugs que ya se corrigieron ──────────────────
print("\n[5] Bugs corregidos que no deben reaparecer")
check("p.filas" not in html, "no quedan referencias a p.filas (bug que rompía el render)")
check("nombres.has(f.proveedor)" in html, "las filas se reconstruyen desde DATA")
check("const r = x => Math.round" in html, "redondeo por fila antes de sumar (cuadratura)")

print("\n" + "=" * 54)
if fails:
    print(f"{len(fails)} FALLA(S):")
    for f in fails:
        print("  -", f)
    sys.exit(1)
print("TODO OK")
